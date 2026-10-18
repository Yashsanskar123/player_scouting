"""Excel exports (openpyxl). Text from users is neutralised against formula injection."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.reports.common import DISCLAIMER, generated_stamp, season_label

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def safe_cell(value):
    """Excel runs text starting with = + - @ as a formula; prefix it so user text stays text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int] | None = None, formats: dict | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([safe_cell(v) for v in row])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths[i - 1] if widths else max(12, min(len(h) + 4, 30)))
    for col, fmt in (formats or {}).items():
        for cell in ws[col][1:]:
            cell.number_format = fmt
    return ws


def _summary(wb: Workbook, title: str, lines: list[tuple[str, object]]):
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = safe_cell(title)
    ws["A1"].font = Font(bold=True, size=14)
    for i, (k, v) in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=safe_cell(v))
    ws.cell(row=len(lines) + 5, column=1, value=DISCLAIMER).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=len(lines) + 5, start_column=1, end_row=len(lines) + 5, end_column=4)
    ws.row_dimensions[len(lines) + 5].height = 45
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 46


def _to_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def shortlist_workbook(shortlist: dict, season: int | None) -> bytes:
    entries = shortlist.get("entries", [])
    wb = Workbook()
    _summary(wb, f"Shortlist: {shortlist.get('name', '')}", [
        ("Players", len(entries)), ("Analysis season", season_label(season) if season else "Latest"),
        ("Generated", generated_stamp()),
    ])
    headers = ["Player", "Position", "Club", "League", "Age", "Status", "Scout rating", "Performance", "Potential",
               "Market value (EUR)", "Fair value (EUR)", "Value gap %", "Risk tier", "Reliability", "Notes", "Added"]
    rows = []
    for e in entries:
        p = e.get("player") or {}
        rows.append([
            p.get("short_name"), p.get("canonical_position"), p.get("club_name"), p.get("league_name"), p.get("age_fifa"),
            e.get("status"), e.get("rating"), p.get("performance_score"), p.get("development_potential"),
            p.get("value_eur"), p.get("fair_value_eur"), p.get("value_gap_pct"), p.get("recruitment_risk_tier"),
            p.get("sample_reliability_tier"), e.get("notes"), (e.get("added_at") or "")[:10],
        ])
    _sheet(wb, "Players", headers, rows, [24, 10, 22, 18, 6, 14, 12, 12, 11, 18, 18, 11, 10, 12, 40, 12],
           {"J": "#,##0", "K": "#,##0", "L": "0.0"})
    return _to_bytes(wb)


def opportunities_workbook(payload: dict, filters: dict) -> bytes:
    results = payload.get("results", [])
    wb = Workbook()
    _summary(wb, "Market opportunities", [
        ("Analysis season", season_label(payload.get("season"))), ("Players exported", len(results)),
        ("Total matching", payload.get("total", len(results))),
        *[(k, v) for k, v in filters.items() if v not in (None, "", "all")],
        ("Generated", generated_stamp()),
        ("Contract data", "Current Transfermarkt snapshot (latest season only)" if (payload.get("contract_context") or {}).get("available") else "Not shown for this season"),
    ])
    headers = ["Player", "Position", "Club", "League", "Age", "Opportunity score", "Performance", "Potential", "Market value (EUR)",
               "Fair value (EUR)", "Value gap %", "Wage/value %", "Risk tier", "Reliability score", "Minutes",
               "Contract status", "Contract expires", "TM value 12m change %"]
    rows = []
    for r in results:
        c = r.get("contract") or {}
        rows.append([
            r.get("short_name"), r.get("canonical_position"), r.get("club_name"), r.get("league_name"), r.get("age_fifa"),
            r.get("opportunity_score"), r.get("performance_score"), r.get("development_potential"), r.get("value_eur"),
            r.get("fair_value_eur"), r.get("value_gap_pct"), r.get("wage_to_value_pct"), r.get("recruitment_risk_tier"),
            r.get("sample_reliability_score"), r.get("playing_time_min"), c.get("contract_status"), c.get("contract_expires"),
            c.get("tm_value_change_12m_pct"),
        ])
    _sheet(wb, "Opportunities", headers, rows, [24, 10, 22, 18, 6, 12, 12, 11, 18, 18, 11, 11, 10, 12, 10, 28, 14, 14],
           {"I": "#,##0", "J": "#,##0", "K": "0.0", "L": "0.00"})
    return _to_bytes(wb)
