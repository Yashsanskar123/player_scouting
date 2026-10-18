from io import BytesIO

import pytest

from openpyxl import load_workbook

from src.reports import common
from src.reports.excel_report import opportunities_workbook, safe_cell, shortlist_workbook
from src.reports.pdf_report import club_report_pdf, player_report_pdf


def test_common_helpers():
    assert common.season_label(2425) == "2024/25" and common.season_label(None) == "—"
    assert common.money(140_500_000) == "€140.5M" and common.money(250_000) == "€250K" and common.money(None) == "—"
    assert common.pdf_safe("Lamine Yamal لامين") .startswith("Lamine Yamal")
    assert common.pdf_safe("Dušan Vlahović") != "—"
    assert common.file_slug("Q4 Targets / Wingers!") == "q4-targets-wingers" and common.file_slug("لامين") == "report"
    assert common.percentile_label("per_90_minutes_prg_p") == "Progressive passes / 90"
    assert common.percentile_label("per_90_minutes_something_new") == "Something new"


def test_player_pdf_is_valid_with_full_and_minimal_inputs():
    full = player_report_pdf(
        {"short_name": "A. Player", "canonical_position": "ST", "performance_score": 71.2, "value_eur": 20e6, "component_attacking": 70,
         "potential_projection_status": "exploratory_empirical", "sample_reliability_tier": "High"},
        {"per_90_minutes_gls": 80.0, "per_90_minutes_clr": 5.0, "per_90_minutes_ast": None},
        [{"season": 2324, "club_name": "Club <A&B>", "playing_time_min": 1800}, {"season": 2425, "club_name": "Club C"}],
        {"available": True, "contract": {"expires": "2027-06-30", "status": "1-2 years left"}, "market": {"current_value_eur": 18e6}, "profile": {}}, 2425)
    assert full[:5] == b"%PDF-" and len(full) > 2500
    assert player_report_pdf({}, None, None, None, None)[:5] == b"%PDF-"  # nothing optional is required
    assert club_report_pdf({"profile": {"club_name": "Arsenal", "formation": "4-3-3"}, "season": 2425})[:5] == b"%PDF-"


def test_excel_neutralises_formulas_and_keeps_data():
    assert safe_cell("=SUM(A1)") == "'=SUM(A1)" and safe_cell("-1+2") == "'-1+2" and safe_cell("ok") == "ok" and safe_cell(5) == 5
    data = {"name": "=evil()", "entries": [{"status": "Tracked", "rating": 4, "notes": "@cmd", "added_at": "2026-10-02T10:00:00",
                                           "player": {"short_name": "P. One", "value_eur": 5_000_000, "club_name": "+Club"}}]}
    wb = load_workbook(BytesIO(shortlist_workbook(data, 2425)))
    assert wb.sheetnames == ["Summary", "Players"]
    ws = wb["Players"]
    assert ws.max_row == 2 and ws["A2"].value == "P. One" and ws["C2"].value == "'+Club" and ws["O2"].value == "'@cmd"
    assert wb["Summary"]["A1"].value == "Shortlist: =evil()"  # starts with text, so Excel will not run it


def test_opportunities_workbook_rows_and_filters():
    payload = {"season": 2425, "total": 2, "contract_context": {"available": True},
               "results": [{"short_name": "X", "value_eur": 1e6, "contract": {"contract_status": "2+ years left", "contract_expires": "2030-06-30"}},
                           {"short_name": "Y", "contract": None}]}
    wb = load_workbook(BytesIO(opportunities_workbook(payload, {"Position": "ST", "Segment": "all", "League": None})))
    assert wb["Opportunities"].max_row == 3 and wb["Opportunities"]["P2"].value == "2+ years left"
    summary = {wb["Summary"].cell(row=r, column=1).value: wb["Summary"].cell(row=r, column=2).value for r in range(3, 14)}
    assert summary["Position"] == "ST" and "League" not in summary and "Segment" not in summary


def test_goalkeeper_report_omits_outfield_measures():
    import subprocess, tempfile, os
    pdf = player_report_pdf({"short_name": "G. Keeper", "canonical_position": "GK", "component_attacking": 98, "component_goalkeeping": 70},
                            {"per_90_minutes_gls": 100.0, "per_90_minutes_performance_saves": 60.0, "per_90_minutes_performance_savepct": 55.0}, [], None, 2425)
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(pdf)
    try:
        try:
            text = subprocess.run(["pdftotext", f.name, "-"], capture_output=True, text=True).stdout
        except FileNotFoundError:
            pytest.skip("pdftotext (poppler) is not installed")
    finally:
        os.unlink(f.name)
    assert "Goalkeeping" in text and "Saves / 90" in text
    assert "Attacking" not in text and "Goals / 90" not in text
